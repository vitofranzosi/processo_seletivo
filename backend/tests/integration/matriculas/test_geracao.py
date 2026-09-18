"""O arquivo de um Edital, sem redigitar nada (`US1`).

**A conferência é no arquivo, e não no código** — é a régua do `C1` do roteiro: 34 cabeçalhos na
grafia do destino, dados a partir da linha 2, uma linha por convocado e nenhuma linha de exemplo.
"""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from processo_seletivo.matriculas.application.exportar import compor, gerar
from processo_seletivo.matriculas.domain import colunas, nomes
from processo_seletivo.matriculas.models import GeracaoDeArquivo

pytestmark = pytest.mark.django_db


def aberto(arquivo):
    """A planilha relida dos bytes entregues — a ida e volta, que é o que prova."""
    return load_workbook(BytesIO(arquivo.conteudo))


def primeira_populacao(edital):
    from processo_seletivo.matriculas.application.populacao import opcoes

    return opcoes(edital)[0]


def gerar_do_cenario(quem_exporta, edital):
    escolhida = primeira_populacao(edital)
    return gerar(
        ator=quem_exporta,
        edital=edital,
        especie=escolhida.especie,
        referencia=escolhida.referencia,
        # **O que a tela teria mostrado.** Recompor aqui é o que o navegador faz ao abrir a prévia;
        # o teste que prova a recusa por resumo obsoleto está em `test_recusas.py`.
        confirmacao_do_resumo=compor(
            ator=quem_exporta,
            edital=edital,
            especie=escolhida.especie,
            referencia=escolhida.referencia,
        ).assinatura,
    )


def test_o_arquivo_tem_uma_aba_com_o_nome_do_destino(cenario, quem_exporta):
    """`FR-436`: uma aba só, chamada `Import_ModeloCefor`."""
    edital, _ = cenario
    livro = aberto(gerar_do_cenario(quem_exporta, edital))
    assert livro.sheetnames == [nomes.ABA]


def test_os_trinta_e_quatro_cabecalhos_saem_na_grafia_do_destino(cenario, quem_exporta):
    """`FR-436`: `A1:AH1`, acentos inclusive.

    Normalizá-los faria o importador não encontrar a coluna.
    """
    edital, _ = cenario
    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    cabecalhos = [celula.value for celula in aba[1]]
    assert cabecalhos == list(colunas.CABECALHOS)
    assert aba["X1"].value == "ENDEREÇO"
    assert aba["Y1"].value == "NÚMERO"
    assert aba["AH1"].value == "COD_POLO"


def test_ha_uma_linha_por_convocado_e_nenhuma_linha_de_exemplo(cenario, quem_exporta):
    """`SC-...`/`C1`: os dados começam na linha 2, e não há linha artificial nenhuma."""
    edital, convocadas = cenario
    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    assert aba.max_row == len(convocadas) + 1
    protocolos = {aba.cell(row=linha, column=1).value for linha in range(2, aba.max_row + 1)}
    assert protocolos == {inscricao.protocolo for inscricao in convocadas}


def test_todas_as_celulas_saem_como_texto(cenario, quem_exporta):
    """`FR-437`: formato `@` em tudo — é o que preserva o zero à esquerda e a data."""
    edital, _ = cenario
    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    for linha in aba.iter_rows():
        for celula in linha:
            assert celula.number_format == "@"
            assert celula.value is None or isinstance(celula.value, str)


def test_a_data_sobrevive_a_ida_e_volta_pelo_excel(cenario, quem_exporta):
    """`SC-144`: `12/07/1994` não vira o serial `34527`."""
    edital, _ = cenario
    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    coluna = colunas.CABECALHOS.index("DATA_NASCIMENTO") + 1
    assert aba.cell(row=2, column=coluna).value == "12/07/1994"


def test_o_zero_a_esquerda_da_zona_eleitoral_sobrevive(cenario, quem_exporta):
    """`SC-152`: `034` continua `034` depois de gravado e relido."""
    edital, _ = cenario
    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    coluna = colunas.CABECALHOS.index("ZONA_ELE") + 1
    assert aba.cell(row=2, column=coluna).value == "034"


def test_o_cep_sai_com_hifen_e_nao_vira_numero(cenario, quem_exporta):
    """`SC-144`: o hífen some assim que a célula deixa de ser texto."""
    edital, _ = cenario
    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    coluna = colunas.CABECALHOS.index("CEP") + 1
    assert aba.cell(row=2, column=coluna).value == "29040-860"


def test_a_ampla_concorrencia_sai_como_ac(cenario, quem_exporta):
    """`FR-443`, `D-003`: derivado da **ausência** de Modalidade."""
    edital, _ = cenario
    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    coluna = colunas.CABECALHOS.index("COD_FORMA_INGRESSO") + 1
    assert {aba.cell(row=linha, column=coluna).value for linha in range(2, aba.max_row + 1)} == {
        "AC"
    }


def test_a_geracao_fica_registrada_com_o_que_a_auditoria_precisa(cenario, quem_exporta):
    """`FR-447`: ator, instante, população, quantidade, versão do resultado e dos mapeamentos."""
    edital, convocadas = cenario
    arquivo = gerar_do_cenario(quem_exporta, edital)
    registro = GeracaoDeArquivo.objects.get(pk=arquivo.geracao.pk)
    assert registro.edital_id == edital.id
    assert registro.gerado_por == quem_exporta.subject
    assert registro.quantidade_de_linhas == len(convocadas)
    assert registro.populacao_especie == nomes.CONVOCACAO
    assert registro.populacao_rotulo.startswith("Convocados")
    assert registro.versao_do_resultado
    assert registro.versao_dos_mapeamentos == colunas.VERSAO_DOS_MAPEAMENTOS


def test_o_registro_guarda_qual_requerimento_estava_vigente(cenario, quem_exporta):
    """§9, *Edge Cases*: é o que responde *"qual era"* depois de uma sucessão.

    Sem guardar uma segunda cópia da declaração — o registro guarda o identificador, e auditar quem
    exportou não pode exigir uma segunda cópia do que foi exportado.
    """
    edital, convocadas = cenario
    arquivo = gerar_do_cenario(quem_exporta, edital)
    registro = GeracaoDeArquivo.objects.get(pk=arquivo.geracao.pk)
    assert len(registro.requerimentos) == len(convocadas)
    assert all(isinstance(identificador, str) for identificador in registro.requerimentos)


def test_o_registro_nao_guarda_dado_de_candidato(cenario, quem_exporta):
    """§18, *Auditoria*: o log que copia o dado sensível **é** o vazamento.

    Nem CPF, nem nome, nem endereço: o registro diz quem exportou, quando e quanto.
    """
    edital, convocadas = cenario
    arquivo = gerar_do_cenario(quem_exporta, edital)
    registro = GeracaoDeArquivo.objects.get(pk=arquivo.geracao.pk)
    texto = " ".join(
        str(valor) for valor in registro.__dict__.values() if not callable(valor)
    ) + " ".join(registro.requerimentos)
    for inscricao in convocadas:
        assert inscricao.cpf_normalizado not in texto
        assert inscricao.nome not in texto
    assert "Rua das Palmeiras" not in texto


def test_o_nome_e_o_email_da_identidade_prevalecem(cenario, quem_exporta):
    """A Inscrição congelou nome e endereço na submissão; a matrícula usa o que vale hoje."""
    from django.utils import timezone

    from processo_seletivo.identidade.application.associacao import associar_credencial
    from processo_seletivo.identidade.models import CandidateIdentity

    edital, convocadas = cenario
    identidade = CandidateIdentity.objects.create(
        subject=convocadas[0].identity_subject,
        nome="Joana Ferreira Gonçalves",
        cpf_normalizado="11144477735",
        created_at=timezone.now(),
    )
    associar_credencial(identidade, "joana@exemplo.br", "Joana@exemplo.br")

    aba = aberto(gerar_do_cenario(quem_exporta, edital)).active
    coluna_nome = colunas.CABECALHOS.index("NOME") + 1
    coluna_email = colunas.CABECALHOS.index("EMAIL") + 1
    linha = next(
        numero
        for numero in range(2, aba.max_row + 1)
        if aba.cell(row=numero, column=1).value == convocadas[0].protocolo
    )
    assert aba.cell(row=linha, column=coluna_nome).value == "Joana Ferreira Gonçalves"
    assert aba.cell(row=linha, column=coluna_email).value == "Joana@exemplo.br"
