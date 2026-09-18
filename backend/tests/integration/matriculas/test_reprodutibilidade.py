"""O mesmo arquivo, duas vezes (`US4`, `FR-446`, `SC-148`).

**Sem isso, duas pessoas gerando no mesmo dia obtêm arquivos diferentes e ninguém sabe qual foi
importado.** A ordem é `CLASSIF_CURSO_FINAL` e, no empate, o protocolo — e como aquela coluna sai
vazia enquanto a `Q-2` não responder, a ordem efetiva é a do protocolo.
"""

from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from processo_seletivo.convocacao.application.selectors import chamada_em_aberto
from processo_seletivo.matriculas.application.exportar import compor, gerar
from processo_seletivo.matriculas.application.populacao import opcoes
from processo_seletivo.matriculas.domain import colunas
from processo_seletivo.publicacoes.application.selectors import effective_version
from processo_seletivo.requerimentos.domain import nomes as requerimento_nomes
from processo_seletivo.requerimentos.models import RequerimentoDeMatricula
from tests.fixtures.matriculas import campos_declarados

pytestmark = [pytest.mark.django_db, pytest.mark.integration]


def gerar_do(quem_exporta, edital):
    escolhida = opcoes(edital)[0]
    return gerar(
        ator=quem_exporta,
        edital=edital,
        especie=escolhida.especie,
        referencia=escolhida.referencia,
        # **O que a tela teria mostrado.** Recompor aqui é o que o navegador faz ao abrir a prévia;
        # o teste que prova a recusa por resumo obsoleto está em `test_recusas.py`.
        confirmacao_do_resumo=compor(
            ator=quem_exporta,
            edital=edital,
            especie=escolhida.especie,
            referencia=escolhida.referencia,
        ).assinatura,
    )


def celulas(arquivo):
    """Todo o conteúdo, célula a célula — que é a régua do `C6`."""
    aba = load_workbook(BytesIO(arquivo.conteudo)).active
    return [[celula.value for celula in linha] for linha in aba.iter_rows()]


def test_duas_geracoes_seguidas_produzem_o_mesmo_conteudo(cenario, quem_exporta):
    """`SC-148`: iguais célula a célula."""
    edital, _ = cenario
    assert celulas(gerar_do(quem_exporta, edital)) == celulas(gerar_do(quem_exporta, edital))


def test_as_linhas_saem_em_ordem_deterministica(cenario, quem_exporta):
    """`FR-446`: sem ordem explícita, duas gerações do mesmo conjunto saem em ordens diferentes."""
    edital, _ = cenario
    protocolos = [linha[0] for linha in celulas(gerar_do(quem_exporta, edital))[1:]]
    assert protocolos == sorted(protocolos)


def test_cada_geracao_deixa_o_seu_registro(cenario, quem_exporta):
    """Gerar de novo não reescreve o registro anterior: a tabela é append-only."""
    from processo_seletivo.matriculas.models import GeracaoDeArquivo

    edital, _ = cenario
    gerar_do(quem_exporta, edital)
    gerar_do(quem_exporta, edital)
    assert GeracaoDeArquivo.objects.filter(edital=edital).count() == 2


def test_o_arquivo_traz_o_requerimento_vigente_e_o_registro_guarda_qual_era(cenario, quem_exporta):
    """§9, *Edge Cases*: sucedido entre duas gerações, o arquivo traz o **vigente**.

    E o registro da primeira continua apontando o requerimento que **estava** vigente naquele dia —
    que é o que torna reconstituível qual declaração foi exportada, sem guardar uma segunda cópia
    dela.
    """
    edital, convocadas = cenario
    primeira = gerar_do(quem_exporta, edital)
    anterior = RequerimentoDeMatricula.objects.get(
        inscricao=convocadas[0], requerimento_anterior__isnull=True
    )

    agora = timezone.now()
    sucessor = RequerimentoDeMatricula.objects.create(
        inscricao=convocadas[0],
        requerimento_anterior=anterior,
        convocacao_autorizadora=chamada_em_aberto(convocadas[0]),
        status=requerimento_nomes.ENVIADO,
        disponibilizado_em=agora,
        created_at=agora,
        enviado_em=agora,
        versao_aceita=effective_version(edital_id=edital.id),
        declaracao_hash="b" * 64,
        declaracao_aceita_em=agora,
        **campos_declarados(bairro="Praia do Canto"),
    )

    segunda = gerar_do(quem_exporta, edital)
    coluna = colunas.CABECALHOS.index("BAIRRO")
    linha = next(linha for linha in celulas(segunda)[1:] if linha[0] == convocadas[0].protocolo)
    assert linha[coluna] == "Praia do Canto"
    assert str(anterior.id) in primeira.geracao.requerimentos
    assert str(sucessor.id) in segunda.geracao.requerimentos
    assert str(sucessor.id) not in primeira.geracao.requerimentos
