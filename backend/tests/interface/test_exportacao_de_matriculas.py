"""A jornada inteira pelo canal de quem conduz, **sem shell** (`031`, Fase 7, `C7`).

**É este arquivo que decide se a feature está entregue.** A Constituição, Princípio VI: *"uma
capacidade que o domínio sustenta mas que nenhuma interface alcança NÃO DEVE ser considerada
entregue"*, e *"demonstrar por chamada manual aquilo que o canal do ator não oferece NÃO satisfaz
esta exigência"*. Tudo aqui passa pelo navegador — sem `manage.py`, sem `psql`, sem chamar função.
"""

from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from processo_seletivo.matriculas.domain import colunas
from processo_seletivo.matriculas.models import GeracaoDeArquivo
from tests.fixtures.matriculas import montar_cenario_da_exportacao
from tests.interface.conftest import identificar

pytestmark = pytest.mark.django_db


@pytest.fixture
def cenario(db, gestor, api_client, manager_headers, process_payload, raiz_de_arquivos):
    """Dois convocados, e o primeiro declarou cor **indígena**.

    É o caso do `R-1`, e ele está aqui porque a prévia que o nomeia é a leitura mais sensível desta
    tela — a que expõe cor/raça de uma pessoa identificada.
    """
    from processo_seletivo.requerimentos.domain import nomes as requerimento_nomes

    return montar_cenario_da_exportacao(
        gestor,
        api_client,
        manager_headers,
        process_payload,
        declaracoes=[{"cor_raca": requerimento_nomes.INDIGENA}],
    )


def entrar(client, papeis=("exportador",)):
    return identificar(client, "exportadora", list(papeis))


def tela(edital, **parametros):
    endereco = reverse("interface:exportar-matriculas", args=[edital.id])
    if parametros:
        from urllib.parse import urlencode

        return f"{endereco}?{urlencode(parametros)}"
    return endereco


def test_a_opcao_aparece_no_menu_do_edital(client, seletor_ligado, cenario):
    """`T031d`: uma capacidade sem porta não é capacidade.

    Quem conduz precisa **encontrar** a exportação no Edital, sem saber a URL de cor.
    """
    edital, _ = cenario
    entrar(client)
    corpo = client.get(reverse("interface:detalhe", args=[edital.id])).content.decode()
    assert "Exportar para matrícula" in corpo
    assert tela(edital) in corpo


def test_o_edital_que_nao_matricula_nao_oferece_a_opcao(
    client,
    seletor_ligado,
    db,
    gestor,
    api_client,
    manager_headers,
    process_payload,
    raiz_de_arquivos,
):
    """`029`, `D-002`: num certame que não coleta requerimento, a capacidade **não existe**.

    Não fica escondida nem desabilitada — não existe. Este sistema conduz Editais de professor
    substituto, de técnico-administrativo, de tutores e de bolsistas, e nenhum deles matricula
    ninguém. Oferecer *Exportar para matrícula* ali seria oferecer um beco.
    """
    from tests.fixtures.matriculas import montar_cenario_da_exportacao
    from tests.fixtures.publicacao import publish_original

    edital, _ = montar_cenario_da_exportacao(
        gestor,
        api_client,
        manager_headers,
        process_payload,
        prefixo="031-tela-sem-requerimento",
        quantos=1,
        publicar=lambda api, cabecalhos, carga, draft: publish_original(
            api, cabecalhos, carga, draft=draft
        ),
    )
    entrar(client)
    corpo = client.get(reverse("interface:detalhe", args=[edital.id])).content.decode()

    assert "Exportar para matrícula" not in corpo


def test_quem_nao_tem_a_permissao_nao_ve_a_opcao(client, seletor_ligado, cenario):
    """`FR-455`: a permissão não é concedida a papel nenhum por padrão.

    O Gestor lê o dossiê de uma pessoa por vez e **não** exporta o conjunto: são atos distintos.
    """
    edital, _ = cenario
    identificar(client, "gestora", ["gestor"])
    corpo = client.get(reverse("interface:detalhe", args=[edital.id])).content.decode()
    assert "Exportar para matrícula" not in corpo


def test_o_endereco_direto_e_recusado_sem_a_permissao(client, seletor_ligado, cenario):
    """`SC-150`: a recusa é do backend, e colar o endereço não a contorna."""
    edital, _ = cenario
    identificar(client, "gestora", ["gestor"])
    assert client.get(tela(edital)).status_code == 403


def test_o_edital_de_outro_escopo_responde_404(client, seletor_ligado, cenario):
    """O identificador público não confere autorização: 404, indistinguível de inexistente."""
    edital, _ = cenario
    identificar(client, "exportadora", ["exportador"], escopo="outro-instituto")
    assert client.get(tela(edital)).status_code == 404


def test_a_tela_exige_escolher_a_populacao(client, seletor_ligado, cenario):
    """`FR-433`: sem escolher, não há o que gerar — e a tela diz que não existe padrão."""
    edital, _ = cenario
    entrar(client)
    corpo = client.get(tela(edital)).content.decode()
    assert "Quem entra no arquivo" in corpo
    assert "Convocados" in corpo
    assert "exportar tudo" in corpo


def test_as_lacunas_aparecem_antes_do_download(client, seletor_ligado, cenario):
    """`UX-060`, `C3`: o resumo vem **antes**, porque depois do download ninguém lê.

    E o aviso da coluna 31 aparece mesmo sem nenhuma outra lacuna (`FR-452`, `SC-153`): se ele não
    estiver lá, o cenário falhou — mesmo com o arquivo perfeito.
    """
    edital, _ = cenario
    entrar(client)
    corpo = client.get(tela(edital, populacao=_primeira(edital))).content.decode()
    assert "O que sairá vazio" in corpo
    assert "COD_CURSO" in corpo and "sistema acadêmico" in corpo
    assert "RENDA_PER_CAPITA_PNP" in corpo
    assert "por pessoa" in corpo
    assert "Baixar o arquivo" in corpo
    # Nada foi gerado: ler o resumo não é exportar.
    assert not GeracaoDeArquivo.objects.exists()


def test_o_download_entrega_a_planilha_e_registra_a_geracao(client, seletor_ligado, cenario):
    """`C7`, passos 6 e 7: o arquivo chega pelo navegador, e o ato fica registrado."""
    edital, convocadas = cenario
    entrar(client)
    resposta = client.post(tela(edital), _formulario_da_previa(client, edital))
    assert resposta.status_code == 200
    assert resposta["Content-Disposition"].startswith("attachment;")
    assert ".xlsx" in resposta["Content-Disposition"]
    # **Não armazenável**: o artefato mais concentrado de dado pessoal não fica no cache.
    assert "no-store" in resposta["Cache-Control"]

    aba = load_workbook(BytesIO(resposta.content)).active
    assert [celula.value for celula in aba[1]] == list(colunas.CABECALHOS)
    assert aba.max_row == len(convocadas) + 1

    registro = GeracaoDeArquivo.objects.get()
    assert registro.gerado_por == "exportadora"
    assert registro.quantidade_de_linhas == len(convocadas)


def test_a_previa_que_nomeia_pessoas_fica_na_trilha(client, seletor_ligado, cenario):
    """Princípio III: ler cor/raça de uma pessoa identificada é acesso a dado sensível.

    A `GeracaoDeArquivo` não alcança este caso de propósito — aquele registro é do **arquivo**, e a
    prévia não gera arquivo nenhum. Sem esta trilha, a leitura mais sensível desta tela seria a
    única que não deixa rastro.

    **E o registro referencia, não copia**: nem o nome da pessoa, nem o que ela declarou.
    """
    from processo_seletivo.auditoria.models import RegistroAuditoria

    edital, convocadas = cenario
    entrar(client)
    corpo = client.get(tela(edital, populacao=_primeira(edital))).content.decode()
    assert "Indígena" in corpo, "a prévia precisa nomear alguém para haver o que registrar"

    registro = RegistroAuditoria.objects.get(operation="MATRICULA_PREVER")
    assert registro.actor_subject == "exportadora"
    assert str(registro.aggregate_id) == str(edital.id)
    assert registro.permission == "matricula:exportar"
    assert "Convocados" in registro.reason
    assert convocadas[0].nome not in registro.reason
    assert "Indígena" not in registro.reason
    # Ler o resumo não é exportar: nenhum arquivo foi gerado.
    assert not GeracaoDeArquivo.objects.exists()


def test_a_previa_sem_ninguem_nomeado_nao_polui_a_trilha(
    client,
    seletor_ligado,
    db,
    gestor,
    api_client,
    manager_headers,
    process_payload,
    raiz_de_arquivos,
):
    """O que motiva o registro é a **exposição nominal**, e não a leitura da tela.

    Registrar toda abertura de página encheria a trilha de ruído e afogaria o caso que importa.
    """
    from processo_seletivo.auditoria.models import RegistroAuditoria

    edital, _ = montar_cenario_da_exportacao(
        gestor, api_client, manager_headers, process_payload, prefixo="031-sem-nominal"
    )
    entrar(client)
    client.get(tela(edital, populacao=_primeira(edital)))

    assert not RegistroAuditoria.objects.filter(operation="MATRICULA_PREVER").exists()


def test_o_download_recusa_quando_o_resumo_envelheceu(client, seletor_ligado, cenario):
    """`UX-060` pelo canal: a assinatura da prévia viaja no formulário, e o POST a confere.

    Aqui o `POST` chega com uma assinatura que já não corresponde — é o que acontece quando alguém
    deixa a página aberta e o mundo anda. A recusa vem para a mesma tela, e nada é gerado.
    """
    edital, _ = cenario
    entrar(client)
    resposta = client.post(
        tela(edital),
        {"populacao": _primeira(edital), "confirmacao_do_resumo": "assinatura-de-outra-leitura"},
    )
    assert resposta.status_code == 409
    assert "não é o que foi conferido" in resposta.content.decode()
    assert not GeracaoDeArquivo.objects.exists()


def test_a_recusa_diz_o_que_falta_e_de_quem_sem_sair_da_tela(
    client,
    seletor_ligado,
    db,
    gestor,
    api_client,
    manager_headers,
    process_payload,
    raiz_de_arquivos,
):
    """`UX-061`: nunca *"não foi possível gerar"*.

    E a recusa volta **para esta tela**, com a escolha à vista: quem conduz decide se cobra a
    pessoa ou se troca de população, sem refazer o caminho.
    """
    from processo_seletivo.convocacao.models import Convocacao
    from processo_seletivo.inscricoes.models import Inscricao
    from tests.fixtures.convocacao import convocar

    edital, _ = montar_cenario_da_exportacao(
        gestor, api_client, manager_headers, process_payload, prefixo="031-tela-recusa", quantos=1
    )
    ja_chamadas = set(
        Convocacao.objects.filter(edital=edital).values_list("inscricao_id", flat=True)
    )
    sem_declarar = (
        Inscricao.objects.filter(edital=edital, status=Inscricao.Status.SUBMETIDA)
        .exclude(pk__in=ja_chamadas)
        .order_by("protocolo")
        .first()
    )
    convocar(edital, gestor, sem_declarar, idempotency_key="031-tela-recusa-extra")

    entrar(client)
    resposta = client.post(tela(edital), _formulario_da_previa(client, edital))
    corpo = resposta.content.decode()
    assert resposta.status_code == 422
    assert "Nada foi gerado" in corpo
    assert sem_declarar.protocolo in corpo
    assert "Quem entra no arquivo" in corpo
    assert not GeracaoDeArquivo.objects.exists()


def _formulario_da_previa(client, edital):
    """O que o navegador envia ao clicar em *Baixar*: a população **e** a assinatura do resumo.

    Lida da própria página, e não recomposta no teste: é assim que se prova que a tela a entrega.
    """
    import re

    corpo = client.get(tela(edital, populacao=_primeira(edital))).content.decode()
    achado = re.search(r'name="confirmacao_do_resumo" value="([^"]*)"', corpo)
    return {
        "populacao": _primeira(edital),
        "confirmacao_do_resumo": achado.group(1) if achado else "",
    }


def _primeira(edital):
    from processo_seletivo.matriculas.application.populacao import opcoes

    escolhida = opcoes(edital)[0]
    return f"{escolhida.especie}:{escolhida.referencia}"
